#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
CleanL10nKeys - Multi-language table health checker.

Runs five checks over the localization Excel tables and their consumers:

  1. unreferenced keys  - keys referenced by NONE of: code (static class field
                          refs + string literals), config data, FGUI customData.
                          These are cleanup candidates (--apply deletes rows).
  2. missing keys       - keys referenced by code (L10nKey./LaunchL10nKey./
                          GetLanguage("x") call sites) or FGUI customData but
                          absent from the tables -> runtime would show "[key]".
  3. translation coverage - per-language empty-cell statistics per table.
  4. FGUI hardcoded text  - static CJK text written directly in FGUI xml that
                          is NOT bound to an L10n key (declared-L10n coverage).
  5. key naming lint    - keys violating ^[a-z][a-z0-9_]*$.

Self-reference exclusions (derived from the tables being cleaned, would make
every key look referenced forever):
  - tblocalization*.json      exported data of the localization tables themselves
  - L10nKey.cs/LaunchL10nKey.cs  generated key classes listing every is_code key

Blind spots: dynamically composed keys ("common_" + x) are statically
undetectable in BOTH directions; config-data-driven keys are only covered in
the unreferenced direction (token set), not in check-missing.

Usage:
  python src/CleanL10nKeys.py          # full health report (all five checks)
  python src/CleanL10nKeys.py --apply  # + delete unreferenced rows from Excel
                                       #   (re-export is manual: gen-client-*.bat)
"""
import re
import sys
import xml.etree.ElementTree as ET
from pathlib import Path

import openpyxl

REPO_ROOT      = Path(__file__).resolve().parents[3]             # repo root (src/ -> tool dir -> Tools -> repo)
CODE_DIR       = REPO_ROOT / 'Unity/Assets/Scripts'              # all runtime code (AOT + Hotfix)
CONFIG_DIR     = REPO_ROOT / 'Unity/Assets/Bundles/Config'       # exported config json
FGUI_DIR       = REPO_ROOT / 'FairyGUIProject/assets'            # FGUI source xml
EXCEL_FILES    = [
    REPO_ROOT / 'Config/Excels/Local/L-Localization-通用.xlsx',
    REPO_ROOT / 'Config/Excels/Local/L-Localization-成就.xlsx',
    REPO_ROOT / 'Config/Excels/Local/L-Localization-设置.xlsx',
    REPO_ROOT / 'Config/Excels/Local/L-LocalizationAOT-aot-热更前.xlsx',
]

TOKEN_RE          = re.compile(r'[A-Za-z0-9_]+')
KEY_RE            = re.compile(r'^[A-Za-z0-9_]+$')
VALID_KEY_RE      = re.compile(r'^[a-z][a-z0-9_]*$')
CJK_RE            = re.compile(r'[㐀-䶿一-鿿]')
SELF_REF_JSON_RE  = re.compile(r'^tblocalization.*\.json$')
SELF_REF_CS       = {'L10nKey.cs', 'LaunchL10nKey.cs'}
# code references resolvable to a concrete key
CODE_REF_RES      = [
    re.compile(r'\bL10nKey\.([A-Za-z0-9_]+)'),
    re.compile(r'\bLaunchL10nKey\.([A-Za-z0-9_]+)'),
    re.compile(r'\bGetLanguage(?:Text)?\(\s*"([A-Za-z0-9_]+)"'),
]
L10N_SEG_RE       = re.compile(r'L10n:([^|]+)')
REPORT_FILE       = Path(__file__).resolve().parents[1] / '多语言配置报告.txt'   # 完整报告输出文件（gitignore）


# ---------------------------------------------------------------- reading ----

def read_table(path):
    """Read one Luban table. Returns (header_names, entries);
    entry = {'row': excel_row, 'key': str, 'values': tuple} with data rows only."""
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    header = [str(c).strip() if c is not None else '' for c in rows[0]]
    key_col = header.index('key')

    entries = []
    for idx, row in enumerate(rows[3:], start=4):   # data starts at Excel row 4
        if key_col >= len(row):
            continue
        raw = row[key_col]
        key = str(raw).strip() if raw is not None else ''
        if not key or not KEY_RE.match(key):
            continue                                 # blank / Luban comment row (#xxx)
        entries.append({'row': idx, 'key': key, 'values': row})
    return header, entries


def table_lang_columns(header):
    """Language value columns = named columns besides structural ones."""
    structural = {'', 'key', 'is_code', '#Description'}
    return [(i, name) for i, name in enumerate(header)
            if name not in structural and not name.startswith('#')]


# ------------------------------------------------------------ reference sources ----

def collect_reference_tokens():
    """Every [A-Za-z0-9_]+ token from all reference sources (for unreferenced check).
    Derived files (self-reference exclusions, see module doc) are silently skipped."""
    tokens = set()

    def feed(path):
        tokens.update(TOKEN_RE.findall(path.read_text(encoding='utf-8', errors='ignore')))

    for path in CODE_DIR.rglob('*.cs'):
        if path.name in SELF_REF_CS:
            continue
        feed(path)
    for path in CONFIG_DIR.glob('*.json'):
        if SELF_REF_JSON_RE.match(path.name):
            continue
        feed(path)
    for path in FGUI_DIR.rglob('*.xml'):
        feed(path)

    return tokens


def parse_l10n_custom_data(custom_data):
    """Parse every key out of a customData string's L10n: segments.
    Handles simple mode (L10n:key) and controller mode (L10n:ctrl,0=k1,1=k2),
    with '|' segment composition."""
    keys = []
    for seg in L10N_SEG_RE.findall(custom_data or ''):
        body = seg.strip()
        if not body:
            continue
        if ',' in body and '=' in body:              # controller mode
            for pair in body.split(',')[1:]:
                kv = pair.split('=')
                if len(kv) == 2 and KEY_RE.match(kv[1].strip()):
                    keys.append(kv[1].strip())
        elif KEY_RE.match(body):
            keys.append(body)
    return keys


def collect_concrete_key_refs():
    """Keys referenced by resolvable sources, for the missing-key check:
    - code: L10nKey.x / LaunchL10nKey.x / GetLanguage("x") call sites
    - FGUI: customData L10n: segments (simple + controller modes)
    Returns set of referenced keys."""
    refs = set()

    for path in CODE_DIR.rglob('*.cs'):
        if path.name in SELF_REF_CS:
            continue
        text = path.read_text(encoding='utf-8', errors='ignore')
        for res in CODE_REF_RES:
            refs.update(res.findall(text))

    for path in FGUI_DIR.rglob('*.xml'):
        try:
            root = ET.parse(path).getroot()
        except ET.ParseError:
            continue
        for elem in root.iter():
            refs.update(parse_l10n_custom_data(elem.get('customData')))

    return refs


# ---------------------------------------------------------------- checks ----

def check_unreferenced(tokens):
    """Returns list of (excel_path, entry) cleanup candidates."""
    all_unreferenced = []
    for excel in EXCEL_FILES:
        if not excel.exists():
            print(f'[跳过] 表不存在：{excel.name}')
            continue
        _, entries = read_table(excel)
        for e in entries:
            if e['key'] not in tokens:
                all_unreferenced.append((excel, e))
    return all_unreferenced


def apply_delete(unreferenced):
    """Delete unreferenced rows per table (bottom-up so indices stay valid).
    Returns report lines describing the deletion (caller owns printing/recording)."""
    by_table = {}
    for excel, e in unreferenced:
        by_table.setdefault(excel, []).append(e)

    lines = []
    deleted = 0
    for excel, entries in by_table.items():
        wb = openpyxl.load_workbook(excel)
        ws = wb.worksheets[0]
        for e in sorted(entries, key=lambda x: x['row'], reverse=True):
            ws.delete_rows(e['row'])
        wb.save(excel)
        deleted += len(entries)
        lines.append(f'[已删除] {excel.name}：{len(entries)} 行')
    lines.append(f'[已删除] 合计 {deleted} 行')
    lines.append('注意：本次仅修改了 Excel。请执行 gen-client-json.bat / gen-client-bin.bat')
    lines.append('      （或 Unity 菜单「FuFramework/配置表/导出配置表」）重新导表。')
    return lines


def check_missing(known_keys):
    """Keys referenced by resolvable sources but absent from every table.
    known_keys: the union of all table keys (computed once in main)."""
    referenced = collect_concrete_key_refs()
    return sorted(referenced - set(known_keys))


def calc_coverage():
    """Per-language empty-cell statistics for every table.
    Returns list of (table_name, total_keys, [(lang, missing_keys)]) for tables with data."""
    tables = []
    for excel in EXCEL_FILES:
        if not excel.exists():
            continue
        header, entries = read_table(excel)
        lang_cols = table_lang_columns(header)
        if not lang_cols or not entries:
            continue

        missing_by_lang = []
        for i, name in lang_cols:
            missing = [e['key'] for e in entries
                       if (e['values'][i] if i < len(e['values']) else None) in (None, '')]
            missing_by_lang.append((name, missing))
        tables.append((excel.name, len(entries), missing_by_lang))
    return tables


def paint(text, has_issue):
    """Color a section summary line: red when issues found, green when clean.
    Unity console understands rich-text color tags; tty consoles get plain text."""
    if sys.stdout.isatty():
        return text
    color = '#ff5050' if has_issue else '#7cff7c'
    return f'<color={color}>{text}</color>'


def scan_fgui_hardcoded():
    """Static CJK text in FGUI xml on components that carry no L10n binding."""
    results = []   # (relative_path, elem_name, text)
    for path in FGUI_DIR.rglob('*.xml'):
        try:
            root = ET.parse(path).getroot()
        except ET.ParseError:
            continue
        rel = path.relative_to(REPO_ROOT)
        for elem in root.iter():
            if parse_l10n_custom_data(elem.get('customData')):
                continue                                  # already bound to L10n
            for attr in ('text', 'title'):
                value = elem.get(attr)
                if value and CJK_RE.search(value):
                    results.append((str(rel), elem.get('name') or elem.tag, value.strip()))
                    break
    return results


def lint_naming():
    """Keys violating ^[a-z][a-z0-9_]*$."""
    violations = []   # (table, key)
    for excel in EXCEL_FILES:
        if not excel.exists():
            continue
        _, entries = read_table(excel)
        for e in entries:
            if not VALID_KEY_RE.match(e['key']):
                violations.append((excel.name, e['key']))
    return violations


# ---------------------------------------------------------------- main ----

def build_report(do_apply):
    """Run all five checks. Returns (summary_lines, full_lines, unreferenced).
    summary_lines: few colored lines for console/dialog; full_lines: complete report."""
    summary = []   # stdout / Unity 对话框：仅摘要（染色）
    full    = []   # 完整报告文件内容（纯文本）

    def head(text, has_issue):
        full.append('-' * 70)
        full.append(text)
        summary.append(paint(text, has_issue))

    full.append('=' * 70)
    full.append(f'多语言表健康检查报告    模式：{"检查 + 清理" if do_apply else "仅检查"}')
    full.append('=' * 70)
    summary.append(f'模式：{"健康检查 + 清理执行" if do_apply else "健康检查"}')

    # ---- 1. 未引用 key（可清理项） ----
    unreferenced = check_unreferenced(collect_reference_tokens())
    head(f'■ 1. 未引用 key（可清理）—— {"发现 %d 个" % len(unreferenced) if unreferenced else "无"}',
         bool(unreferenced))
    if not unreferenced:
        full.append('    所有 key 均被引用，无需清理。')
    else:
        by_table = {}
        for excel, e in unreferenced:
            by_table.setdefault(excel.name, []).append(e['key'])
        for table_name, keys in by_table.items():
            full.append(f'    【{table_name}】未引用 {len(keys)} 个：')
            full.extend(f'        {key}' for key in keys)

    # ---- 2. 缺失 key（引用了但表里没有） ----
    all_keys = {e['key'] for p in EXCEL_FILES if p.exists() for _, es in [read_table(p)] for e in es}
    missing = check_missing(all_keys)
    head(f'■ 2. 缺失 key（已被引用但表中不存在）—— {"发现 %d 个" % len(missing) if missing else "无缺失"}',
         bool(missing))
    if not missing:
        full.append('    所有引用的 key 都在表中。')
    else:
        full.extend(f'    {key}' for key in missing)

    # ---- 3. 翻译覆盖率 ----
    coverage = calc_coverage()
    coverage_bad = any(m for _, _, langs in coverage for _, m in langs)
    head('■ 3. 翻译覆盖率 —— ' + ('存在未填写语言' if coverage_bad else '全部已填写'), coverage_bad)
    for table_name, total, langs in coverage:
        full.append(f'    【{table_name}】共 {total} 个 key：')
        for name, missing_list in langs:
            if not missing_list:
                full.append(f'        {name}: 全部已填写')
            else:
                preview = '、'.join(missing_list[:5]) + ('...' if len(missing_list) > 5 else '')
                full.append(f'        {name}: 缺 {len(missing_list)}/{total}（{preview}）')

    # ---- 4. FGUI 硬编码文本 ----
    hardcoded = scan_fgui_hardcoded()
    head(f'■ 4. FGUI 硬编码文本（未绑定 L10n 的静态中文）—— {"发现 %d 处" % len(hardcoded) if hardcoded else "未发现"}',
         bool(hardcoded))
    if not hardcoded:
        full.append('    FGUI 静态文本均已声明式绑定或无中文。')
    else:
        full.extend(f'    {rel}  [{name}]  "{text[:30]}"' for rel, name, text in hardcoded)

    # ---- 5. key 命名规范 ----
    violations = lint_naming()
    head(f'■ 5. key 命名规范（要求 ^[a-z][a-z0-9_]*$）—— {"违规 %d 个" % len(violations) if violations else "全部合规"}',
         bool(violations))
    if not violations:
        full.append('    全部合规。')
    else:
        full.extend(f'    {table}: {key}' for table, key in violations)

    full.append('=' * 70)
    summary_line = (f'汇总：未引用 {len(unreferenced)} | 缺失 {len(missing)} | '
                    f'硬编码 {len(hardcoded)} | 命名违规 {len(violations)}')
    full.append(summary_line)
    summary.append(summary_line)
    summary.append(f'完整报告：{REPORT_FILE}')
    return summary, full, unreferenced


def main():
    # Console (tty) keeps the system code page (GBK cmd windows); piped output
    # (Unity editor captures stdout) switches to UTF-8 as both ends agreed.
    if not sys.stdout.isatty():
        sys.stdout.reconfigure(encoding='utf-8', errors='replace')
    else:
        sys.stdout.reconfigure(errors='replace')
    do_apply = '--apply' in sys.argv[1:]

    summary, full, unreferenced = build_report(do_apply)

    # ---- 清理执行（仅未引用行，导表手动） ----
    if do_apply:
        if unreferenced:
            deleted = apply_delete(unreferenced)
            summary.extend(deleted)
            full.extend(deleted)
        else:
            full.append('没有可清理项。')

    REPORT_FILE.write_text('\n'.join(full) + '\n', encoding='utf-8')

    # stdout 只有摘要与文件路径，报告本体在文件中（项目膨胀后报告可达数百行，控制台不承载）
    print('\n'.join(summary))
    if not do_apply and unreferenced:
        print('当前为预览，未做任何修改。加 --apply 参数（或 Unity 菜单「多语言健康检查—清理」）删除未引用行。')


if __name__ == '__main__':
    main()
